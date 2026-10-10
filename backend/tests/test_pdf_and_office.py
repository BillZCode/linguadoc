"""
Tests for PDF and Office (XLSX, PPTX) Document Processors.
"""

import os
import pytest
import pymupdf
import openpyxl

from models.translation_job import TranslationJob, TranslationConfig, TranslationProvider, ToneStyle
from pipeline.pipeline import TranslationPipeline

FIXTURES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
TEST_OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_outputs")


@pytest.fixture
def pipeline():
    return TranslationPipeline(output_dir=TEST_OUTPUT_DIR)


def test_pdf_processor(pipeline):
    # Generate simple test PDF
    pdf_path = os.path.join(FIXTURES_DIR, "test_document.pdf")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Chapter 1: The Beginning", fontsize=16)
    page.insert_text((50, 120), "This is the first paragraph of the PDF document.", fontsize=11)
    page.insert_text((50, 160), "「Are we ready to proceed?」 asked the engineer.", fontsize=11)
    doc.save(pdf_path)
    doc.close()
    
    config = TranslationConfig(
        source_language="en",
        target_language="id",
        provider=TranslationProvider.DEMO
    )
    job = TranslationJob(
        original_filename="test_document.pdf",
        file_format="pdf",
        file_size_bytes=os.path.getsize(pdf_path),
        source_file_path=pdf_path,
        config=config
    )
    
    completed_job = pipeline.run_pipeline(job)
    assert completed_job.stage.value == "ready"
    assert os.path.exists(completed_job.output_file_path)
    
    # Verify reconstructed PDF can be opened by pymupdf
    out_doc = pymupdf.open(completed_job.output_file_path)
    assert len(out_doc) == 1
    out_text = out_doc[0].get_text()
    assert len(out_text.strip()) > 0
    out_doc.close()


def test_xlsx_formula_protection(pipeline):
    # Generate test XLSX with formulas and numbers
    xlsx_path = os.path.join(FIXTURES_DIR, "test_financial.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Summary"
    
    ws["A1"] = "Item Description"
    ws["B1"] = "Quantity"
    ws["C1"] = "Unit Price"
    ws["D1"] = "Total Cost"
    
    ws["A2"] = "Server Cloud Hosting"
    ws["B2"] = 5
    ws["C2"] = 120.50
    ws["D2"] = "=B2*C2"  # Formula!
    
    wb.save(xlsx_path)
    
    config = TranslationConfig(
        source_language="en",
        target_language="id",
        provider=TranslationProvider.DEMO
    )
    job = TranslationJob(
        original_filename="test_financial.xlsx",
        file_format="xlsx",
        file_size_bytes=os.path.getsize(xlsx_path),
        source_file_path=xlsx_path,
        config=config
    )
    
    completed_job = pipeline.run_pipeline(job)
    assert completed_job.stage.value == "ready"
    assert os.path.exists(completed_job.output_file_path)
    
    # Reopen and verify formula was NOT translated or corrupted!
    out_wb = openpyxl.load_workbook(completed_job.output_file_path, data_only=False)
    out_ws = out_wb["Financial Summary"]
    assert out_ws["D2"].value == "=B2*C2"  # Formula preserved!
    assert out_ws["B2"].value == 5         # Numeric integer preserved!
    assert out_ws["C2"].value == 120.50    # Float preserved!
